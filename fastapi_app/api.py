# --- Standardbibliotheken ---
import base64
import json
import os
import re
import tempfile
import traceback
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional

# --- FastAPI & Middleware ---
from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

# --- Pydantic & Validation ---
from pydantic import BaseModel, Field, model_validator, ConfigDict
from jsonschema import validate, ValidationError

# --- PDF & ReportLab ---
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas
from PyPDF2 import PdfReader  # bevorzugt statt pypdf.PdfReader

# --- DOCX Verarbeitung ---
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

# --- Eigene Module ---
from analysis.analysis import readability, bloom_level
from analysis.answer_comparator import AnswerComparator
from analysis.readability_module import analyze_german_readability, ReadabilityAnalyzer
from mistral import MistralWrapper

# --- Server ---
import uvicorn





# ----------------- ExamAssistantChat UI --------------------
# Create FastAPI app
app = FastAPI(title="My API Base", version="1.0.0")
wrapper = MistralWrapper()
comparator = AnswerComparator()
readability_analyzer = ReadabilityAnalyzer()

# ============================================================
# CORS (Frontend-Port 8001)
# ============================================================
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://127.0.0.1:8001", "http://localhost:8001"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class OpenApiMessage(BaseModel):
    content: str

class Attachment(BaseModel):
    file_name: str
    content_base64: str
    file_type: str = "pdf"

class MessageRequest(BaseModel):
    chat_id: str                    
    content: str
    attachment: Optional[Attachment] = None




@app.post("/sendApiMessage")
def sendApiMessage(message: MessageRequest):
    pass
    
  
        







# ============================================================
# Hilfsfunktionen , Comparator
# ============================================================

#llm extraction
def now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def append_jsonl(path: Path, obj: dict) -> None:
    ensure_dir(path.parent)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")    

        
def ensure_dir(p: Path) -> Path:
    p.mkdir(parents=True, exist_ok=True)
    return p


def generate_doc_id(name: str) -> str:
    base = os.path.splitext(os.path.basename(name))[0][:50] or "upload"
    return f"{base}_{uuid.uuid4().hex[:8]}"


def make_run_id() -> str:
    ts = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    return f"{ts}_{uuid.uuid4().hex[:8]}"


def pdf_to_text(path: str, max_chars=10000) -> str:
    """Extrahiert Text aus einer PDF-Datei."""
    text = ""
    try:
        reader = PdfReader(path)
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text += content + "\n"
    except Exception:
        pass
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_chars] if text else ""


def export_solution_as_pdf(text: str, export_path: Path):
    """Speichert generierten Lösungstext als PDF."""
    export_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(export_path), pagesize=A4)
    width, height = A4
    x, y = 50, height - 50
    max_width = width - 100
    for line in text.split("\n"):
        words = line.split(" ")
        current_line = ""
        for w in words:
            test_line = current_line + w + " "
            if c.stringWidth(test_line, "Helvetica", 10) < max_width:
                current_line = test_line
            else:
                c.drawString(x, y, current_line)
                y -= 14
                current_line = w + " "
                if y < 60:
                    c.showPage()
                    y = height - 50
        if current_line:
            c.drawString(x, y, current_line)
            y -= 14
            if y < 60:
                c.showPage()
                y = height - 50
    c.save()

# ============================================================
# API: Lösung generieren
# ============================================================
@app.post("/generate-solution")
async def generate_solution(chat_id: str = Form("default"), pdf: Optional[UploadFile] = File(None)):
    try:
        if not pdf:
            raise HTTPException(status_code=400, detail="Keine PDF-Datei hochgeladen.")
        raw = await pdf.read()
        if not raw.startswith(b"%PDF"):
            raise HTTPException(status_code=400, detail="Ungültige PDF-Datei.")

        # === Basisverzeichnis vorbereiten ===
        doc_id = generate_doc_id(pdf.filename)
        base_dir = Path("data") / "uploads" / chat_id / doc_id
        ensure_dir(base_dir)
        pdf_path = base_dir / "source.pdf"
        pdf_path.write_bytes(raw)

        # === Text aus PDF extrahieren ===
        text = pdf_to_text(str(pdf_path))
        if not text:
            raise HTTPException(status_code=400, detail="PDF enthält keinen lesbaren Text.")
        # 🔧 Fix 1 – zusätzliche Bereinigung für konsistente Regex-Erkennung
        text = re.sub(r"\s+", " ", text).strip()

        # === Aufgabenblöcke erkennen ===
        aufgaben_raw = re.findall(
            r"(Aufgabe\s*\d+[a-zA-Z]?[.:–-]?\s.*?)(?=(?:Aufgabe\s*\d+[a-zA-Z]?[.:–-]?|Frage\s*\d+[.:–-]?|$))",
            text,
            flags=re.IGNORECASE | re.DOTALL,
        )

        aufgaben = []
        for a in aufgaben_raw:
            clean = re.sub(r"^(?:Aufgabe|Frage)\s*\d+[a-zA-Z]?[.:–-]?\s*", "", a.strip(), flags=re.IGNORECASE)
            aufgaben.append(clean.strip())

        if not aufgaben:
            aufgaben = [text]

        # === Prompt dynamisch aufbauen ===
        prompt_parts = []
        for i, frage in enumerate(aufgaben, start=1):
            prompt_parts.append(f"Aufgabe {i}:\n{frage}\n\nAntwort {i}:")
        structured_prompt = "\n\n".join(prompt_parts)

        prompt = (
            "Beantworte präzise und sachlich nur die folgenden Aufgaben.\n"
            "Formatiere das Ergebnis **genau** in diesem Stil:\n\n"
            "Aufgabe 1:\n[Fragetext]\n\nAntwort 1:\n[Lösung]\n\n"
            "Aufgabe 2:\n[Fragetext]\n\nAntwort 2:\n[Lösung]\n\n"
            "usw.\n\n"
            "Gib **nur** die Fragen und Antworten zurück — keine Einleitung, keine Kommentare.\n\n"
            f"Hier sind die Aufgaben:\n<<<\n{structured_prompt}\n>>>"
        )

        run_id = make_run_id()
        llm_dir = Path("data") / "llm" / chat_id / run_id
        ensure_dir(llm_dir)
        (llm_dir / "prompt.txt").write_text(prompt, encoding="utf-8")

        try:
            reply = wrapper.send_request(prompt)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"LLM Fehler: {str(e)}")

        if "Service tier capacity exceeded" in str(reply):
            reply = "⚠️ Der LLM-Dienst ist derzeit ausgelastet. Bitte versuche es später erneut."

        # === Speichern & PDF-Export ===
        (llm_dir / "response.txt").write_text(reply, encoding="utf-8")
        export_dir = Path("exports") / "solutions" / chat_id
        export_path = export_dir / f"{run_id}.pdf"
        export_solution_as_pdf(reply, export_path)

        return {
            "chat_id": chat_id,
            "doc_id": doc_id,
            "run_id": run_id,
            "solution": reply,
            "export_pdf": str(export_path),
            "aufgaben_gefunden": len(aufgaben),
        }

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Fehler in generate_solution: {e}")


# ============================================================
# API: Vergleichsanalyse
# ============================================================
@app.post("/compare-solutions")
async def compare_solutions(
    file1: UploadFile = File(...),
    file2: UploadFile = File(...),
    mode: str = Form("quick"),
    chat_id: str = Form("default"),
):
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as t1:
            t1.write(await file1.read())
            path_ml = t1.name
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as t2:
            t2.write(await file2.read())
            path_sl = t2.name

        model_text = pdf_to_text(path_ml)
        student_text = pdf_to_text(path_sl)
        if not model_text or not student_text:
            raise HTTPException(status_code=400, detail="Leere oder unlesbare PDFs.")

        run_id = make_run_id()
        comp = AnswerComparator()

        # === QUICK ===
        if mode == "quick":
            quick_score = comp.quick_compare(student_text, model_text)
            readability_score = analyze_german_readability(student_text)
            return {
                "mode": "quick",
                "quick_similarity": quick_score,
                "readability": readability_score,
                "summary": f"Quick-Vergleich: {round(quick_score*100,2)}% Ähnlichkeit erkannt."
            }

        # === DETAILED ===
        elif mode == "detailed":
            summary = comp.detailed_compare(student_text, model_text)
            if not summary or not isinstance(summary, dict):
                raise ValueError("Keine gültige summary-Daten erhalten.")

            # 🔧 Fix 2 – Lesbarkeitsanalyse der Musterlösung ergänzen
            if "readability_metrics_model" not in summary:
                summary["readability_metrics_model"] = analyze_german_readability(model_text)

            # Struktur-Fallbacks
            if "tasks" not in summary or not summary["tasks"]:
                summary["tasks"] = [{
                    "task_id": "Gesamtanalyse",
                    "similarity": summary.get("similarity_score", 0.0),
                    "model_solution": model_text or "(Keine Musterlösung gefunden)",
                    "student_answer": student_text or "(Keine Studentenlösung gefunden)",
                    "feedback": summary.get("feedback", "(Kein Feedback verfügbar)")
                }]
            else:
                for t in summary["tasks"]:
                    t["task_id"] = t.get("task_id", "Unbekannte Aufgabe")
                    t["model_solution"] = t.get("model_solution") or "(Keine Musterlösung gefunden)"
                    t["student_answer"] = t.get("student_answer") or "(Keine Studentenlösung gefunden)"
                    t["feedback"] = t.get("feedback") or "(Kein Feedback verfügbar)"

            outdir = Path("analysis") / chat_id / run_id
            outdir.mkdir(parents=True, exist_ok=True)
            with open(outdir / "summary.json", "w", encoding="utf-8") as f:
                json.dump(summary, f, indent=2, ensure_ascii=False)

            return {
                "mode": "detailed",
                "chat_id": chat_id,
                "run_id": run_id,
                "summary": summary,
            }

        raise HTTPException(status_code=400, detail=f"Ungültiger Modus: {mode}")

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Fehler in compare_solutions: {e}")


# ============================================================
# API: Detaillierte Analyse abrufen
# ============================================================
@app.get("/detailed-analysis")
async def detailed_analysis(chat_id: str = Query(...), run_id: str = Query(...)):
    try:
        summary_path = Path("analysis") / chat_id / run_id / "summary.json"
        if not summary_path.exists():
            raise HTTPException(status_code=404, detail="Analyse-Datei nicht gefunden.")

        with open(summary_path, "r", encoding="utf-8") as f:
            summary = json.load(f)

        # Struktur garantieren (Fallback)
        if "tasks" not in summary or not summary["tasks"]:
            summary["tasks"] = [{
                "task_id": "Gesamtanalyse",
                "similarity": summary.get("similarity_score", 0.0),
                "model_solution": summary.get("model_solution", "(Keine Musterlösung gefunden)"),
                "student_answer": summary.get("student_answer", "(Keine Studentenlösung gefunden)"),
                "feedback": summary.get("feedback", "(Kein Feedback verfügbar)")
            }]
        else:
            for t in summary["tasks"]:
                t["task_id"] = t.get("task_id", "Unbekannte Aufgabe")
                t["model_solution"] = t.get("model_solution") or "(Keine Musterlösung gefunden)"
                t["student_answer"] = t.get("student_answer") or "(Keine Studentenlösung gefunden)"
                t["feedback"] = t.get("feedback") or "(Kein Feedback verfügbar)"

        return {"chat_id": chat_id, "run_id": run_id, "summary": summary}

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Fehler beim Laden der Analyse: {e}")





# -------- AufgabenGenerator ---------

# --- Modelle ---


class ChatIdOnly(BaseModel):
    chat_id: str
    exam_id: str


class TopicItem(BaseModel):
    name: str
    weight: Optional[int] = 0

class ExportBody(BaseModel):
    chat_id: str
    exam_id: str
    title: Optional[str] = "Klausur"
    subject: Optional[str] = None
    topics: Optional[List[TopicItem]] = []


class TaskPrompt(BaseModel):
    chat_id: str
    prompt: str
    schema_type: str = "default"
    retry_of: Optional[str] = None
    
class AcceptBody(BaseModel):
    # Ignoriere unerwartete Felder (z. B. falls ein alter Client noch "json" mitsendet)
    model_config = ConfigDict(extra='ignore')

    chat_id: str
    exam_id: str
    run_id: str
    text: str
    payload: Optional[dict] = None
    schema_type: Optional[str] = "default"

    @model_validator(mode="before")
    @classmethod
    def merge_json_to_payload(cls, data):
        # robust gegen None / falsche Typen
        if isinstance(data, dict):
            # wenn nur "json" kommt -> auf payload mappen
            if 'payload' not in data and 'json' in data:
                data['payload'] = data.pop('json')
            # wenn BEIDE kommen -> payload hat Vorrang; "json" ignorieren
            elif 'payload' in data and 'json' in data:
                data.pop('json', None)
        return data





    
def exam_dir(chat_id: str, exam_id: str) -> Path:
    return Path("data") / "exports" / chat_id / exam_id

def accepted_dir(chat_id: str) -> Path:
    return Path("data") / "exports" / chat_id

def accepted_list_path(chat_id: str, exam_id: str) -> Path:
    return exam_dir(chat_id, exam_id) / "accepted.json"

def load_accepted(chat_id: str, exam_id: str) -> list:
    p = accepted_list_path(chat_id, exam_id)
    if not p.exists(): return []
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return []
    


from pathlib import Path

def load_prompt(path_env: str, fallback: str) -> str:
    p = os.getenv(path_env)
    if p and Path(p).exists():
        return Path(p).read_text(encoding="utf-8")
    return fallback    

def save_accepted(chat_id: str, exam_id: str, items: list) -> None:
    d = exam_dir(chat_id, exam_id)
    d.mkdir(parents=True, exist_ok=True)
    (d / "accepted.jsonl").open("a", encoding="utf-8").write(
        json.dumps({"ts": now_iso(), "count": len(items)}, ensure_ascii=False) + "\n"
    )
    accepted_list_path(chat_id, exam_id).write_text(
        json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def extract_json_block(text: str) -> dict:
    fenced = re.search(r"```json(.*?)```", text, re.S|re.I)
    payload = fenced.group(1).strip() if fenced else text.strip()
    start = payload.find("{")
    end   = payload.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("No JSON object found in LLM output.")
    raw = payload[start:end+1]
    obj = json.loads(raw)
    if not isinstance(obj, dict):
        raise ValueError(f"Expected JSON object, got {type(obj).__name__}")
    return obj

def render_compact_preview(kind: str, obj: dict) -> str:
    """
    Kompakte Mensch-lesbare Vorschau fürs UI & PDF (aus JSON).
    Unterstützt:
      - default:  { title: str, task: str, ... }
      - mc:       { title, question, options[...] }
      - calc:     { title, problem }
    """
    if not isinstance(obj, dict):
        return str(obj)  # Fallback

    if kind == "mc":
        lines = []
        title = obj.get("title", "Multiple-Choice-Aufgabe")
        q     = obj.get("question", "")
        opts  = obj.get("options", []) or []
        lines.append(f"{title}")
        if q:
            lines.append(q)
        for i, opt in enumerate(opts, 1):
            t = (opt or {}).get("text", "")
            lines.append(f"{i}) {t}")
        return "\n".join([ln for ln in lines if ln]).strip()

    if kind == "calc":
        title = obj.get("title", "Rechenaufgabe")
        prob  = obj.get("problem", "")
        return "\n".join([title, prob]).strip()

    # ---- default (Freitext) ----
    title = obj.get("title", "Aufgabe")
    task_field = obj.get("task")

    # task kann String ODER Objekt sein (abwärtskompatibel)
    if isinstance(task_field, str):
        text = task_field.strip()
    elif isinstance(task_field, dict):
        text = (task_field.get("text") or "").strip()
    else:
        text = ""

    # optionales Fallback – falls du altformat task_text noch irgendwo hast
    if not text:
        text = (obj.get("task_text") or "").strip()

    lines = [title]
    if text:
        lines.append(text)
    return "\n".join(lines).strip()



def group_items_by_topic(items: list, topics_from_body: Optional[List[TopicItem]]):
    def item_topic(it):
        j = (it.get("json") or {}) if isinstance(it.get("json"), dict) else {}
        t = (j.get("topic") or "").strip()
        return t if t else "Allgemein"

    grouped = {}
    for it in items:
        t = item_topic(it)
        grouped.setdefault(t, []).append(it)

    order_from_body = [t.name for t in (topics_from_body or [])]
    remaining = [t for t in grouped.keys() if t not in order_from_body]
    topic_order = order_from_body + sorted(remaining)
    return grouped, topic_order

#------ENDPOINTS-------
@app.post("/task/reset_exam")
def task_reset_exam(body: ChatIdOnly):
    d = exam_dir(body.chat_id, body.exam_id)
    d.mkdir(parents=True, exist_ok=True)
    (d / "accepted.json").write_text("[]", encoding="utf-8")
    (d / "accepted.jsonl").open("a", encoding="utf-8").write(
        json.dumps({"ts": now_iso(), "event": "reset"}, ensure_ascii=False) + "\n"
    )
    return {"ok": True}


@app.post("/task/generate")
def task_generate(body: TaskPrompt):
    # 1) Schema-Typ bestimmen
    schema_key = (body.schema_type or "default").lower()
    if schema_key not in {"default", "mc", "calc"}:
        schema_key = "default"

    # 2) System-Prompt + Schema-Datei wählen
    if schema_key == "mc":
        system_prompt = load_prompt(
            "SYSTEM_PROMPT_MC_PATH",
            "Du bist ein Aufgaben-Generator für Multiple-Choice-Aufgaben. "
            "Formuliere kurz, klar, prüfungsgeeignet."
        )
        schema_path = Path("schemas/multiple_choice_task.schema.json")
        kind = "mc"
    elif schema_key == "calc":
        system_prompt = load_prompt(
            "SYSTEM_PROMPT_CALC_PATH",
            "Du bist ein Aufgaben-Generator für Rechenaufgaben. "
            "Liefer präzise Aufgabenstellung und Lösungsidee (optional), kompakt."
        )
        schema_path = Path("schemas/calculation_task.schema.json")
        kind = "calc"
    else:
        system_prompt = load_prompt(
            "SYSTEM_PROMPT_DEFAULT_PATH",
            "Du bist ein Aufgaben-Generator für offene Prüfungsaufgaben. "
            "Kompakt, ohne Floskeln, nur die Aufgabe."
        )
        schema_path = Path("schemas/default_task.schema.json")
        kind = "default"

    # 3) Schema laden
    if not schema_path.exists():
        raise HTTPException(status_code=500, detail=f"Schema file not found: {schema_path}")
    schema_json = json.loads(schema_path.read_text(encoding="utf-8"))

    # 4) finalen Prompt bauen
    user_prompt = (body.prompt or "").strip() or "Erstelle eine verständliche Prüfungsaufgabe."
    instructions = (
        "Liefere NUR ein einziges JSON-Objekt, das exakt dem folgenden JSON-Schema entspricht. "
        "Kein Fließtext, keine Erklärungen, KEINE Markdown-Codeblöcke außerhalb des JSONs."
    )
    full_prompt = (
        f"{system_prompt}\n\n"
        f"{instructions}\n\n"
        f"JSON-Schema:\n{json.dumps(schema_json, ensure_ascii=False, indent=2)}\n\n"
        f"Benutzerwunsch:\n{user_prompt}"
    )

    # 5) LLM-Run-Ordner
    run_id = make_run_id()
    llm_run_dir = Path("data") / "llm" / body.chat_id / run_id
    llm_run_dir.mkdir(parents=True, exist_ok=True)

    prompt_path = llm_run_dir / "prompt.txt"
    prompt_path.write_text(full_prompt, encoding="utf-8")

    # 6) LLM call
    raw_text = wrapper.send_request(full_prompt)

    # 7) JSON extrahieren & validieren
    try:
        obj = extract_json_block(raw_text)
        validate(instance=obj, schema=schema_json)
    except (ValueError, ValidationError, json.JSONDecodeError) as e:
        (llm_run_dir / "response_raw.txt").write_text(raw_text, encoding="utf-8")
        raise HTTPException(status_code=400, detail=f"Invalid LLM output: {e}")

    # 8) Kompakte Vorschau + Lesbarkeit
    preview_text = render_compact_preview(kind, obj)
    (llm_run_dir / "response.json").write_text(
        json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (llm_run_dir / "rendered.txt").write_text(preview_text, encoding="utf-8")

    rb = readability(preview_text)
    bloom = bloom_level(preview_text)

    # 9) Log
    chat_log_path = Path("data") / "chats" / body.chat_id / "log.jsonl"
    append_jsonl(chat_log_path, {
        "ts": now_iso(),
        "run_id": run_id,
        "mode": "generator",
        "schema": schema_key,
        "prompt": user_prompt,
        "reply_json": obj,
        "preview": preview_text[:500],
        "readability": rb
    })

    # 10) Antwort an UI
    return {
        "run_id": run_id,
        "schema": schema_key,
        "json": obj,
        "text": preview_text,
        "readability": rb,
        "bloom": bloom,  # NEU
        "prompt_path": str(prompt_path),
        "llm_run_dir": str(llm_run_dir),
        "rendered_path": str(llm_run_dir / "rendered.txt"),
        "response_path": str(llm_run_dir / "response.json"),
    }





@app.post("/task/accept")
def task_accept(body: AcceptBody):
    if not body.exam_id:
        raise HTTPException(status_code=400, detail="exam_id fehlt.")
    items = load_accepted(body.chat_id, body.exam_id)

    payload = body.payload or {}
    if isinstance(payload, dict):
        # Fallback-Topic, wenn keins gesetzt wurde (UI setzt es normalerweise)
        topic = (payload.get("topic") or "").strip()
        if not topic:
            payload["topic"] = "Allgemein"

    items.append({
        "run_id": body.run_id,
        "text": body.text,
        "json": payload,
        "schema_type": (body.schema_type or "default"),
        "accepted_at": now_iso()
    })
    save_accepted(body.chat_id, body.exam_id, items)
    return {"ok": True, "count": len(items)}


@app.post("/task/accepted_list")
def task_accepted_list(body: ChatIdOnly):
    items = load_accepted(body.chat_id, body.exam_id)
    return {"ok": True, "items": items}



@app.post("/task/export_pdf")
def task_export_pdf(body: ExportBody):
    items = load_accepted(body.chat_id, body.exam_id)
    if not items:
        raise HTTPException(status_code=400, detail="Keine angenommenen Aufgaben vorhanden.")

    out_dir = exam_dir(body.chat_id, body.exam_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf"

    c = canvas.Canvas(str(out_path), pagesize=A4)
    width, height = A4

    def draw_wrapped(text, x, y, max_width, leading=16, font="Helvetica", size=11):
        c.setFont(font, size)
        lines = simpleSplit(text, font, size, max_width)
        for line in lines:
            if y < 60:
                c.showPage()
                y = height - 50
                c.setFont(font, size)
            c.drawString(x, y, line)
            y -= leading
        return y

    # -------- Deckblatt --------
    y = height - 80
    c.setFont("Helvetica-Bold", 22)
    title = body.subject or body.title or "Klausur"
    c.drawString(40, y, f"Klausur: {title}")
    y -= 32

    c.setFont("Helvetica", 12)
    y = draw_wrapped("Name: _________________________________", 40, y, width-80, leading=20)
    y = draw_wrapped("Vorname: ______________________________", 40, y, width-80, leading=20)
    y = draw_wrapped("Matrikelnummer: _______________________", 40, y, width-80, leading=20)
    y -= 10

    c.drawString(40, y, f"Erstellt am: {datetime.utcnow().strftime('%d.%m.%Y %H:%M UTC')}")
    y -= 30

    if body.topics:
        c.setFont("Helvetica-Bold", 14)
        c.drawString(40, y, "Themen & Gewichtungen")
        y -= 20
        c.setFont("Helvetica", 11)
        for t in body.topics:
            line = f"– {t.name}" + (f" ({t.weight} Punkte)" if (t.weight or 0) > 0 else "")
            y = draw_wrapped(line, 40, y, width-80, leading=16)

    c.showPage()  # Aufgaben beginnen auf neuer Seite

    # -------- Items nach Thema gruppieren --------
    # 1) Topic aus Item ziehen (json.topic) – Fallback "Allgemein"
    def item_topic(it):
        j = (it.get("json") or {}) if isinstance(it.get("json"), dict) else {}
        t = (j.get("topic") or "").strip()
        return t if t else "Allgemein"

    grouped = {}
    for it in items:
        t = item_topic(it)
        grouped.setdefault(t, []).append(it)

    # 2) Reihenfolge der Themen:
    #    zuerst wie vom Wizard übergeben (body.topics), danach evtl. übrige
    order_from_body = [t.name for t in (body.topics or [])]
    remaining = [t for t in grouped.keys() if t not in order_from_body]
    topic_order = order_from_body + sorted(remaining)

    # -------- Druck-Layout --------
    margin_x = 40
    max_w = width - 80

    for topic_name in topic_order:
        topic_items = grouped.get(topic_name, [])
        if not topic_items:
            continue

        # Jede Themen-Gruppe auf einer neuen Seite beginnen
        c.setFont("Helvetica-Bold", 16)
        c.drawString(margin_x, height - 60, f"Thema: {topic_name}")
        # Gewicht, falls in body.topics enthalten
        tcfg = next((t for t in (body.topics or []) if t.name == topic_name), None)
        if tcfg and (tcfg.weight or 0) > 0:
            c.setFont("Helvetica", 11)
            c.drawString(margin_x, height - 80, f"Gewichtung: {tcfg.weight} Punkte")

        y = height - 110
        task_no = 1  # Nummerierung je Thema neu anfangen

        for it in topic_items:
            # Überschrift der Aufgabe
            c.setFont("Helvetica-Bold", 14)
            c.drawString(margin_x, y, f"Aufgabe {task_no}")
            y -= 22

            # Aufgabe-Text (kompakte Vorschau aus Generator)
            c.setFont("Helvetica", 11)
            text = it.get("text") or ""
            y = draw_wrapped(text, margin_x, y, max_w, leading=16)

            # Antwort-Linien NUR für Freitext & Rechenaufgaben
            schema = (it.get("schema_type") or "default").lower()
            if schema in ("default", "calc"):   # KEINE Linien für "mc"
                y -= 8
                for _ in range(6):
                    if y < 70:
                        c.showPage()
                        # Themenkopf NICHT erneut zeichnen, nur weiterlaufen
                        y = height - 50
                        c.setFont("Helvetica", 11)
                    c.line(margin_x, y, width - 40, y)
                    y -= 16
                y -= 10
            else:
                # MC: etwas Abstand, aber keine Linien
                y -= 16

            # Falls Seite voll, neue Seite
            if y < 100:
                c.showPage()
                y = height - 50
                c.setFont("Helvetica", 11)

            task_no += 1

        # Nach Thema: immer Seitenumbruch (nächste Themen-Seite sauber starten)
        c.showPage()

    # Wenn am Ende eine leere Seite entstanden ist (z. B. kein weiteres Thema),
    # ist das okay; ReportLab lässt sie einfach stehen. Optional könnte man
    # noch "c._pagesize" o. Ä. prüfen, aber nicht nötig.
    c.save()
    return {"ok": True, "pdf_path": str(out_path)}



@app.post("/task/export_docx")
def task_export_docx(body: ExportBody):
    items = load_accepted(body.chat_id, body.exam_id)
    if not items:
        raise HTTPException(status_code=400, detail="Keine angenommenen Aufgaben vorhanden.")

    out_dir = exam_dir(body.chat_id, body.exam_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.docx"

    doc = Document()

    # ---------- Deckblatt ----------
    title = body.subject or body.title or "Klausur"
    h = doc.add_heading(f"Klausur: {title}", level=1)
    for p in [ "Name: _________________________________",
               "Vorname: ______________________________",
               "Matrikelnummer: _______________________" ]:
        para = doc.add_paragraph(p)
        para.runs[0].font.size = Pt(12)

    doc.add_paragraph(f"Erstellt am: {datetime.utcnow().strftime('%d.%m.%Y %H:%M UTC')}")

    if body.topics:
        doc.add_heading("Themen & Gewichtungen", level=2)
        for t in body.topics:
            line = f"– {t.name}" + (f" ({t.weight} Punkte)" if (t.weight or 0) > 0 else "")
            doc.add_paragraph(line)

    # Seitenumbruch vor den Aufgaben
    doc.add_page_break()

    # ---------- Gruppierung nach Thema ----------
    grouped, topic_order = group_items_by_topic(items, body.topics)

    def add_answer_lines(n=6):
        for _ in range(n):
            # „Zeile“: einfach ein Unterstrich-Run, gut editierbar
            run = doc.add_paragraph().add_run("_" * 100)
            run.font.size = Pt(11)

    # ---------- Themen ausgeben ----------
    for topic_name in topic_order:
        topic_items = grouped.get(topic_name, [])
        if not topic_items:
            continue

        doc.add_heading(f"Thema: {topic_name}", level=2)
        tcfg = next((t for t in (body.topics or []) if t.name == topic_name), None)
        if tcfg and (tcfg.weight or 0) > 0:
            doc.add_paragraph(f"Gewichtung: {tcfg.weight} Punkte")

        task_no = 1
        for it in topic_items:
            # Überschrift
            doc.add_heading(f"Aufgabe {task_no}", level=3)
            # Aufgabe-Text
            p = doc.add_paragraph(it.get("text") or "")
            p.runs[0].font.size = Pt(11)

            # Platz für Antworten NUR für default & calc
            schema = (it.get("schema_type") or "default").lower()
            if schema in ("default", "calc"):
                add_answer_lines(6)
            else:
                # MC: nur etwas Abstand
                doc.add_paragraph("")

            task_no += 1

        # Jede Themen-Gruppe auf neuer Seite
        doc.add_page_break()

    # speichern
    doc.save(str(out_path))
    return {"ok": True, "docx_path": str(out_path)}


# ============================================================
# Main
# ============================================================
if __name__ == "__main__":
    uvicorn.run("api:app", host="0.0.0.0", port=8000, reload=True)


    # ============================================================
# Root
# ============================================================
@app.get("/")
def index():
    return {"status": "ExamAssistant API aktiv", "version": "3.3.0"}
